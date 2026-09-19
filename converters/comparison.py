"""Bounded, persisted before/after summaries for completed conversions."""
from __future__ import annotations

import csv
import difflib
import json
import os
from pathlib import Path
from typing import Any


MAX_TEXT_CHARS = 12_000
MAX_PREVIEW_ROWS = 5
MAX_COMPARISON_INPUT_BYTES = int(os.getenv('STREAMDOCK_MAX_COMPARISON_INPUT_BYTES', str(16 * 1024 * 1024)))
TABLE_FORMATS = {'csv', 'tsv', 'json', 'ndjson', 'xlsx'}
TEXT_FORMATS = {'txt', 'md', 'markdown', 'html', 'rtf', 'docx'}
IMAGE_FORMATS = {'png', 'jpg', 'jpeg', 'webp', 'bmp', 'tiff', 'gif', 'ico'}
MEDIA_FORMATS = {'mp3', 'wav', 'm4a', 'aac', 'flac', 'ogg', 'opus', 'aiff', 'wma', 'amr', 'mp4', 'mov', 'mkv', 'webm', 'avi', 'flv', 'm4v', '3gp', 'ts'}


def _text_content(path: Path, format_name: str) -> str:
    if path.stat().st_size > MAX_COMPARISON_INPUT_BYTES:
        raise RuntimeError(f'文件超过前后对比读取上限 {MAX_COMPARISON_INPUT_BYTES} 字节')
    if format_name == 'docx':
        from docx import Document  # type: ignore
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        seen = set()
        for section in document.sections:
            for container in (section.header, section.footer):
                key = str(container.part.partname)
                if key in seen:
                    continue
                seen.add(key); parts.extend(paragraph.text for paragraph in container.paragraphs)
        parts.extend(f'[批注 {comment.author}] {comment.text}' for comment in getattr(document, 'comments', []))
        return '\n'.join(part for part in parts if part)
    raw = path.read_text(encoding='utf-8', errors='replace')
    if format_name == 'rtf':
        from .adapters.document_basic import _rtf_to_text
        return _rtf_to_text(raw)
    if format_name == 'html':
        from .adapters.document_basic import _html_to_text
        return _html_to_text(raw)
    return raw


def _text_summary(path: Path, format_name: str) -> dict[str, Any]:
    content = _text_content(path, format_name)
    sample = content[:MAX_TEXT_CHARS]
    return {
        'characters': len(content),
        'lines': len(content.splitlines()),
        'sample': sample,
        'truncated': len(content) > len(sample),
    }


def _table_summary(path: Path, format_name: str) -> dict[str, Any]:
    rows: list[list[Any]]
    structure = {'sheets': 1, 'sheetNames': [], 'formulas': 0, 'mergedRanges': 0, 'hiddenRows': 0, 'hiddenColumns': 0, 'charts': 0}
    if format_name in {'csv', 'tsv'}:
        with path.open('r', encoding='utf-8-sig', newline='') as handle:
            reader = csv.reader(handle, delimiter='\t' if format_name == 'tsv' else ',')
            header = next(reader, [])
            preview_rows = []
            row_count = 0
            for row in reader:
                row_count += 1
                if len(preview_rows) < MAX_PREVIEW_ROWS:
                    preview_rows.append(row)
            rows = [header, *preview_rows]
    elif format_name == 'json':
        if path.stat().st_size > MAX_COMPARISON_INPUT_BYTES:
            raise RuntimeError(f'JSON 超过前后对比读取上限 {MAX_COMPARISON_INPUT_BYTES} 字节')
        payload = json.loads(path.read_text(encoding='utf-8'))
        items = payload if isinstance(payload, list) else [payload]
        keys = sorted({key for item in items if isinstance(item, dict) for key in item})
        rows = [keys] + [[item.get(key) if isinstance(item, dict) else item for key in keys] for item in items]
    elif format_name == 'ndjson':
        if path.stat().st_size > MAX_COMPARISON_INPUT_BYTES:
            raise RuntimeError(f'NDJSON 超过前后对比读取上限 {MAX_COMPARISON_INPUT_BYTES} 字节')
        items = [json.loads(line) for line in path.read_text(encoding='utf-8').splitlines() if line.strip()]
        keys = sorted({key for item in items if isinstance(item, dict) for key in item})
        rows = [keys] + [[item.get(key) if isinstance(item, dict) else item for key in keys] for item in items]
    else:
        if path.stat().st_size > MAX_COMPARISON_INPUT_BYTES:
            raise RuntimeError(f'工作簿超过前后对比读取上限 {MAX_COMPARISON_INPUT_BYTES} 字节')
        from openpyxl import load_workbook  # type: ignore
        workbook = load_workbook(path, read_only=False, data_only=False)
        structure = {
            'sheets': len(workbook.sheetnames),
            'sheetNames': list(workbook.sheetnames),
            'formulas': sum(1 for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith('=')),
            'mergedRanges': sum(len(sheet.merged_cells.ranges) for sheet in workbook.worksheets),
            'hiddenRows': sum(sum(1 for item in sheet.row_dimensions.values() if item.hidden) for sheet in workbook.worksheets),
            'hiddenColumns': sum(sum(1 for item in sheet.column_dimensions.values() if item.hidden) for sheet in workbook.worksheets),
            'charts': sum(len(sheet._charts) for sheet in workbook.worksheets),
        }
        sheet = workbook[workbook.sheetnames[0]]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    headers = [str(value) if value is not None else '' for value in (rows[0] if rows else [])]
    preview = [[str(value) if value is not None else '' for value in row] for row in rows[1:1 + MAX_PREVIEW_ROWS]]
    return {'rows': row_count if format_name in {'csv', 'tsv'} else max(0, len(rows) - 1), 'columns': len(headers), 'headers': headers, 'preview': preview, **structure}


def _image_summary(path: Path) -> dict[str, Any]:
    from PIL import Image  # type: ignore
    with Image.open(path) as image:
        return {
            'width': image.width,
            'height': image.height,
            'mode': image.mode,
            'frames': int(getattr(image, 'n_frames', 1) or 1),
            'hasTransparency': 'A' in image.getbands() or 'transparency' in image.info,
            'iccProfile': bool(image.info.get('icc_profile')),
            'exif': bool(image.getexif()),
        }


def _media_summary(path: Path) -> dict[str, Any]:
    from runtime_checks import validate_media_output
    report = validate_media_output(path)
    return {key: report.get(key) for key in ('durationSeconds', 'hasVideo', 'hasAudio', 'width', 'height', 'videoCodec', 'audioCodec', 'bitrate', 'format')}


def build_conversion_comparison(source: str, target: str, input_path: Path, output_path: Path) -> dict[str, Any]:
    """Return a small comparison artifact without failing an otherwise valid conversion."""
    source, target = source.lower(), target.lower()
    try:
        if source in TABLE_FORMATS and target in TABLE_FORMATS:
            before, after = _table_summary(input_path, source), _table_summary(output_path, target)
            changed_preview_cells = sum(
                1
                for row_index in range(max(len(before['preview']), len(after['preview'])))
                for column_index in range(max(before['columns'], after['columns']))
                if (before['preview'][row_index][column_index] if row_index < len(before['preview']) and column_index < len(before['preview'][row_index]) else None)
                != (after['preview'][row_index][column_index] if row_index < len(after['preview']) and column_index < len(after['preview'][row_index]) else None)
            )
            return {
                'available': True,
                'kind': 'table',
                'before': before,
                'after': after,
                'schemaChanged': before['headers'] != after['headers'],
                'rowDelta': after['rows'] - before['rows'],
                'columnDelta': after['columns'] - before['columns'],
                'changedPreviewCells': changed_preview_cells,
                'workbookStructureChanged': any(before.get(key) != after.get(key) for key in ('sheets', 'sheetNames', 'formulas', 'mergedRanges', 'hiddenRows', 'hiddenColumns', 'charts')),
            }
        if source in TEXT_FORMATS and target in TEXT_FORMATS:
            before, after = _text_summary(input_path, source), _text_summary(output_path, target)
            diff = list(difflib.unified_diff(before['sample'].splitlines(), after['sample'].splitlines(), fromfile='转换前', tofile='转换后', lineterm=''))[:120]
            added = sum(1 for line in diff if line.startswith('+') and not line.startswith('+++'))
            removed = sum(1 for line in diff if line.startswith('-') and not line.startswith('---'))
            return {'available': True, 'kind': 'text', 'before': before, 'after': after, 'diff': diff, 'diffTruncated': len(diff) >= 120, 'addedLines': added, 'removedLines': removed}
        if source in IMAGE_FORMATS and target in IMAGE_FORMATS:
            return {'available': True, 'kind': 'image', 'before': _image_summary(input_path), 'after': _image_summary(output_path)}
        if source in MEDIA_FORMATS and target in MEDIA_FORMATS:
            return {'available': True, 'kind': 'media', 'before': _media_summary(input_path), 'after': _media_summary(output_path)}
        return {'available': False, 'reason': '该转换路径暂未提供结构化前后对比'}
    except Exception as exc:
        return {'available': False, 'reason': f'生成前后对比失败：{exc}'}
