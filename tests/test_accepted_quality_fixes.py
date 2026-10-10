"""Disk-backed regressions; real HTTP/engine acceptance is a separate script."""
import csv
import json
import tempfile
import unittest
from pathlib import Path
from converters.adapters.data import convert_data, read_xlsx_rows
from converters.adapters.document_basic import _html_to_text, _html_to_markdown
from converters.adapters.subtitle import txt_to_srt
from fetchers.subtitle_asr import normalize_asr_boundaries, validate_asr_cues, SubtitleQualityError
from fetchers.subtitle_ocr import OcrSubtitleCue, cues_to_srt
from subtitles.service import parse_subtitles

class AcceptedQualityFixes(unittest.TestCase):
    def test_xlsx_blank_first_sheet_duplicate_headers_and_positions(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);w=Workbook();w.active['Z100'].number_format='0.00';s=w.create_sheet('数据')
            s.append(['姓名',None,'姓名']);s.append(['中文😀','中间','不覆盖'])
            source=root/'complex.xlsx';w.save(source)
            for target,delimiter in [('csv',','),('tsv','\t')]:
                out=root/f'out.{target}';convert_data('xlsx',target,source,out)
                with out.open(encoding='utf-8',newline='') as f:
                    self.assertEqual(list(csv.reader(f,delimiter=delimiter)),[['姓名','','姓名'],['中文😀','中间','不覆盖']])
            with self.assertRaisesRegex(RuntimeError,'表头重复'):read_xlsx_rows(source)

    def test_order_and_literal_strings_roundtrip(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'in.json';out=root/'out.xlsx'
            source.write_text(json.dumps([{'z':'=1+1','a':'#N/A','中文':'😀'}],ensure_ascii=False))
            convert_data('json','xlsx',source,out)
            w=load_workbook(out);self.assertEqual([c.value for c in w.active[1]],['z','a','中文'])
            self.assertTrue(all(c.data_type=='s' for c in w.active[2]));w.close()
            out2=root/'out.json';convert_data('xlsx','json',out,out2)
            self.assertEqual(json.loads(out2.read_text()),json.loads(source.read_text()))

    def test_xml_arbitrary_keys_roundtrip(self):
        import xml.etree.ElementTree as ET
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'in.json';out=root/'out.xml';back=root/'back.json'
            data={'业务 字段':'中文😀','1key':{'nested':'value'},'entry':'ordinary'}
            source.write_text(json.dumps(data,ensure_ascii=False));convert_data('json','xml',source,out)
            ET.parse(out);convert_data('xml','json',out,back)
            self.assertEqual(json.loads(back.read_text())['root'],data)

    def test_html_boundaries_and_noise(self):
        text='<head><title>noise</title></head><script>bad()</script><p>第一段😀</p><table><tr><td>收入</td><td>120</td></tr><tr><td>成本</td><td>20</td></tr></table>'
        result=_html_to_text(text);self.assertIn('收入\t120',result);self.assertIn('\n成本\t20',result)
        self.assertNotIn('noise',result);self.assertNotIn('bad()',_html_to_markdown(text))

    def test_long_txt_and_short_srt_cues(self):
        s=txt_to_srt('\n'.join(f'中文😀{i}' for i in range(3661)))
        d=parse_subtitles(s,filename='long.srt');self.assertEqual(len(d.cues),3661)
        self.assertIn('01:01:00,000 --> 01:01:01,000',s)
        cues=[OcrSubtitleCue(0,.18,'第一句'),OcrSubtitleCue(.18,.4,'第二句')]
        parsed=parse_subtitles(cues_to_srt(cues),filename='short.srt')
        self.assertEqual(parsed.cues[0].end,.18)

    def test_asr_jitter_preserves_text_and_rejects_large_overlap(self):
        cues=[OcrSubtitleCue(0,1.1,'第一句'),OcrSubtitleCue(1,2,'第二句')]
        fixed=normalize_asr_boundaries(cues);validate_asr_cues(fixed)
        self.assertEqual(fixed[0].end,1);self.assertEqual([c.text for c in fixed],[c.text for c in cues])
        with self.assertRaises(SubtitleQualityError):validate_asr_cues(normalize_asr_boundaries([OcrSubtitleCue(0,1.5,'一'),OcrSubtitleCue(1,2,'二')]))

    def test_subtitle_bom_and_gb18030(self):
        from app import decode_subtitle_bytes
        for enc in ('utf-8-sig','utf-16','gb18030'):
            text='中文测试\n字幕';self.assertEqual(decode_subtitle_bytes(text.encode(enc)),text)
