#!/usr/bin/env python3
"""Run every executable non-PDF conversion route with generated real files."""
from __future__ import annotations

import argparse
import bz2
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tarfile
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from converters.models import ConversionLevel  # noqa: E402
from converters.pipeline import convert_file  # noqa: E402
from converters.registry import list_capabilities  # noqa: E402
from runtime_checks import augmented_path, validate_media_output  # noqa: E402
from scripts.test_conversion_robustness import generate_fixtures  # noqa: E402

os.environ['PATH'] = augmented_path()
CONTENT_MARKERS = ('StreamDock', '张三', '中文', '第一', '你好', '标题', '演示文稿')
RELEASE_CONTRACT = ROOT / 'scripts' / 'conversion_release_contract.json'


def command(args: list[str], **kwargs) -> None:
    subprocess.run(args, check=True, capture_output=True, env={**os.environ, 'PATH': augmented_path()}, **kwargs)


def require_content_marker(text: str) -> None:
    if not any(marker in text for marker in CONTENT_MARKERS):
        raise AssertionError('key semantic content marker missing')


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def artifact_summary(path: Path) -> dict[str, Any]:
    if path.is_file():
        return {'type': 'file', 'sha256': file_sha256(path), 'bytes': path.stat().st_size}
    members = [
        {'path': str(item.relative_to(path)), 'sha256': file_sha256(item), 'bytes': item.stat().st_size}
        for item in sorted(path.rglob('*')) if item.is_file()
    ]
    return {'type': 'directory', 'files': len(members), 'bytes': sum(item['bytes'] for item in members), 'members': members}


def independent_rtf_text(path: Path) -> str:
    with tempfile.TemporaryDirectory(prefix='streamdock-matrix-rtf-') as temp_dir:
        root = Path(temp_dir); source = root / 'source.rtf'; shutil.copy2(path, source)
        profile = root / 'profile'; profile.mkdir()
        command(['soffice', '--headless', f'-env:UserInstallation={profile.resolve().as_uri()}', '--convert-to', 'txt:Text', '--outdir', str(root), str(source)])
        output = root / 'source.txt'
        if not output.is_file():
            raise AssertionError('LibreOffice did not reopen RTF output')
        return output.read_text(encoding='utf-8', errors='strict')


def build_all_fixtures(workdir: Path) -> dict[str, Path]:
    fixtures = generate_fixtures(workdir)
    root = workdir / 'fixtures'
    fixtures['json_object'] = root / 'object.json'
    fixtures['json_object'].write_text(json.dumps({'title': 'StreamDock 中文 😀', 'nested': {'rtl': 'مرحبا', 'count': 3}}, ensure_ascii=False), encoding='utf-8')
    from PIL import Image

    with Image.open(fixtures['png']) as image:
        rgb = image.convert('RGB')
        for name, fmt, mode in (
            ('bmp', 'BMP', 'RGB'), ('webp', 'WEBP', 'RGBA'), ('tiff', 'TIFF', 'RGBA'),
            ('ico', 'ICO', 'RGBA'), ('ppm', 'PPM', 'RGB'), ('pgm', 'PPM', 'L'), ('pbm', 'PPM', '1'),
        ):
            path = root / f'image.{name}'; image.convert(mode).save(path, format=fmt); fixtures[name] = path
        fixtures['jpeg'] = root / 'image.jpeg'; rgb.save(fixtures['jpeg'], format='JPEG', quality=94)
        fixtures['pnm'] = root / 'image.pnm'; shutil.copy2(fixtures['ppm'], fixtures['pnm'])

    fixtures['markdown'] = fixtures['md']
    fixtures['tar'] = root / 'sample.tar'
    with tarfile.open(fixtures['tar'], 'w') as archive:
        archive.add(fixtures['folder'], arcname='folder_src')
    fixtures['bz2'] = root / 'plain.txt.bz2'
    with bz2.open(fixtures['bz2'], 'wb') as handle:
        handle.write(fixtures['txt'].read_bytes())

    ffmpeg = shutil.which('ffmpeg', path=augmented_path())
    if ffmpeg:
        wav = fixtures['wav']
        audio_commands = {
            'mp3': ['-i', str(wav), '-c:a', 'libmp3lame'],
            'm4a': ['-i', str(wav), '-c:a', 'aac'],
            'aac': ['-i', str(wav), '-c:a', 'aac', '-f', 'adts'],
            'flac': ['-i', str(wav), '-c:a', 'flac'],
            'ogg': ['-i', str(wav), '-ac', '2', '-c:a', 'vorbis', '-strict', '-2'],
            'opus': ['-i', str(wav), '-c:a', 'libopus'],
            'aiff': ['-i', str(wav), '-c:a', 'pcm_s16be'],
            'wma': ['-i', str(wav), '-c:a', 'wmav2'],
        }
        for fmt, options in audio_commands.items():
            path = root / f'audio.{fmt}'; command([ffmpeg, '-y', '-hide_banner', '-loglevel', 'error', *options, str(path)]); fixtures[fmt] = path
        video = fixtures['mp4']
        video_commands = {
            'mov': ['-i', str(video), '-c', 'copy'],
            'mkv': ['-i', str(video), '-c', 'copy'],
            'webm': ['-i', str(video), '-c:v', 'libvpx-vp9', '-c:a', 'libopus'],
            'avi': ['-i', str(video), '-c:v', 'mpeg4', '-c:a', 'libmp3lame'],
            'flv': ['-i', str(video), '-c:v', 'flv', '-c:a', 'aac'],
            'm4v': ['-i', str(video), '-an', '-c:v', 'mpeg4', '-f', 'm4v'],
            '3gp': ['-i', str(video), '-vf', 'scale=176:144', '-c:v', 'h263', '-c:a', 'aac', '-ar', '8000', '-ac', '1'],
            'ts': ['-i', str(video), '-c', 'copy', '-f', 'mpegts'],
        }
        for fmt, options in video_commands.items():
            path = root / f'video.{fmt}'; command([ffmpeg, '-y', '-hide_banner', '-loglevel', 'error', *options, str(path)]); fixtures[fmt] = path

    try:
        from docx import Document
        docx = root / 'office.docx'; document = Document(); document.add_heading('StreamDock 真实 Office 文档', 1); document.add_paragraph('中文 emoji 😀 RTL مرحبا'); document.save(docx); fixtures['docx'] = docx
        soffice = shutil.which('soffice', path=augmented_path())
        if soffice:
            office_dir = root / 'office-formats'; office_dir.mkdir()
            for fmt in ('doc', 'odt'):
                command([soffice, '--headless', '--convert-to', fmt, '--outdir', str(office_dir), str(docx)])
                fixtures[fmt] = office_dir / f'office.{fmt}'
            from pptx import Presentation
            presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text = 'StreamDock 演示文稿 😀'
            slide.placeholders[1].text = '中文内容 RTL مرحبا'
            pptx = root / 'slides.pptx'; presentation.save(pptx); fixtures['pptx'] = pptx
            for fmt, export_filter in (('ppt', 'ppt:MS PowerPoint 97'), ('odp', 'odp:impress8')):
                command([soffice, '--headless', '--convert-to', export_filter, '--outdir', str(office_dir), str(pptx)])
                converted = office_dir / f'slides.{fmt}'
                if not converted.is_file():
                    raise RuntimeError(f'LibreOffice did not create the real {fmt.upper()} fixture')
                fixtures[fmt] = converted
    except Exception:
        pass

    try:
        from ebooklib import epub
        book = epub.EpubBook(); book.set_identifier('streamdock-real'); book.set_title('StreamDock 电子书'); book.set_language('zh-CN')
        chapter = epub.EpubHtml(title='第一章', file_name='chapter.xhtml', lang='zh-CN'); chapter.content = '<h1>第一章</h1><p>StreamDock 中文 emoji 😀 RTL مرحبا</p>'
        book.add_item(chapter); book.toc = (chapter,); book.spine = ['nav', chapter]; book.add_item(epub.EpubNcx()); book.add_item(epub.EpubNav())
        fixtures['epub'] = root / 'book.epub'; epub.write_epub(str(fixtures['epub']), book)
    except Exception:
        pass

    fixtures['svg'] = root / 'drawing.svg'
    fixtures['svg'].write_text('<svg xmlns="http://www.w3.org/2000/svg" width="160" height="90"><rect width="160" height="90" fill="#2457d6"/><circle cx="80" cy="45" r="25" fill="#ffd43b"/></svg>', encoding='utf-8')
    return fixtures


def validate_output(source: str, target: str, output: Path) -> dict[str, Any]:
    if output.is_dir():
        files = [path for path in output.rglob('*') if path.is_file()]
        if not files: raise AssertionError('empty output directory')
        if source == 'gif' and target == 'png':
            from PIL import Image
            for frame in files:
                with Image.open(frame) as image: image.load()
        return {'files': len(files), 'bytes': sum(path.stat().st_size for path in files)}
    if target == 'folder':
        files = [path for path in output.rglob('*') if path.is_file()]
        if not files:
            raise AssertionError('empty extracted folder')
        return {'files': len(files), 'bytes': sum(path.stat().st_size for path in files)}
    if target == 'zip':
        if not zipfile.is_zipfile(output): raise AssertionError('invalid zip')
        with zipfile.ZipFile(output) as archive:
            bad = archive.testzip()
            if bad: raise AssertionError(f'zip CRC failure: {bad}')
            return {'members': len(archive.infolist())}
    if target in {'tar', 'tar.gz'}:
        if not tarfile.is_tarfile(output): raise AssertionError('invalid tar')
        return {'members': len(tarfile.open(output).getmembers())}
    if target in {'png', 'jpg', 'jpeg', 'webp', 'bmp', 'tiff', 'gif', 'ico', 'ppm', 'pgm', 'pbm', 'pnm'}:
        from PIL import Image, ImageStat
        with Image.open(output) as image:
            image.seek(0); image.load(); extrema = ImageStat.Stat(image.convert('RGB')).extrema
            if not any(high > low for low, high in extrema): raise AssertionError('blank/solid image')
            return {'format': image.format, 'size': image.size, 'frames': int(getattr(image, 'n_frames', 1))}
    if target in {'mp3', 'wav', 'm4a', 'aac', 'flac', 'ogg', 'opus', 'mp4', 'mov', 'mkv', 'webm', 'avi', 'flv', 'm4v', '3gp', 'ts'}:
        expected = 'video' if target in {'mp4', 'mov', 'mkv', 'webm', 'avi', 'flv', 'm4v', '3gp', 'ts'} else 'audio'
        report = validate_media_output(output, expected_kind=expected)
        command([shutil.which('ffmpeg', path=augmented_path()) or 'ffmpeg', '-v', 'error', '-i', str(output), '-f', 'null', '-'])
        ffmpeg = shutil.which('ffmpeg', path=augmented_path()) or 'ffmpeg'
        if report.get('hasAudio') and source != 'amr':
            volume = subprocess.run([ffmpeg, '-hide_banner', '-i', str(output), '-af', 'volumedetect', '-f', 'null', '-'], text=True, capture_output=True, timeout=30)
            import re
            match = re.search(r'max_volume:\s*(-?[\d.]+) dB', volume.stderr)
            if not match or float(match.group(1)) <= -80:
                raise AssertionError('silent or unreadable audio signal')
        if report.get('hasVideo'):
            frame = subprocess.run([ffmpeg, '-v', 'error', '-i', str(output), '-frames:v', '1', '-f', 'image2pipe', '-vcodec', 'png', '-'], capture_output=True, timeout=30, check=True)
            from PIL import Image, ImageStat
            import io
            with Image.open(io.BytesIO(frame.stdout)) as image:
                if not any(high > low for low, high in ImageStat.Stat(image.convert('RGB')).extrema):
                    raise AssertionError('black/solid video frame')
        return report
    if target == 'json':
        raw = output.read_text(encoding='utf-8'); value = json.loads(raw); require_content_marker(raw); return {'type': type(value).__name__}
    if target == 'xml':
        import xml.etree.ElementTree as ET
        raw = output.read_text(encoding='utf-8'); require_content_marker(raw); return {'root': ET.parse(output).getroot().tag}
    if target == 'yaml':
        import yaml
        raw = output.read_text(encoding='utf-8'); require_content_marker(raw); return {'type': type(yaml.safe_load(raw)).__name__}
    if target == 'toml':
        import tomllib
        raw = output.read_text(encoding='utf-8'); require_content_marker(raw); return {'keys': len(tomllib.loads(raw))}
    if target in {'csv', 'tsv', 'txt', 'md', 'markdown', 'html', 'rtf', 'srt', 'vtt'}:
        text = output.read_text(encoding='utf-8')
        if not text.strip() or '\ufffd' in text: raise AssertionError('empty or replacement-character text')
        if target == 'rtf':
            require_content_marker(independent_rtf_text(output))
        else:
            require_content_marker(text)
        if target in {'srt', 'vtt'} and '-->' not in text: raise AssertionError('subtitle timeline missing')
        return {'characters': len(text)}
    if target == 'xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(output, read_only=False, data_only=False); values = [cell.value for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None]
        result = {'sheets': len(workbook.sheetnames), 'values': len(values)}; workbook.close()
        if not values: raise AssertionError('empty workbook')
        require_content_marker('\n'.join(str(value) for value in values))
        return result
    if target == 'docx':
        from docx import Document
        document = Document(output); text = '\n'.join([paragraph.text for paragraph in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        if not text.strip(): raise AssertionError('empty docx')
        require_content_marker(text)
        return {'characters': len(text), 'tables': len(document.tables)}
    if target == 'pptx':
        from pptx import Presentation
        presentation = Presentation(output); text = '\n'.join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, 'text'))
        if not text.strip(): raise AssertionError('empty pptx')
        require_content_marker(text)
        return {'slides': len(presentation.slides), 'characters': len(text)}
    raise AssertionError(f'no validator for target {target}')


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('--workdir', type=Path)
    parser.add_argument('--keep', action='store_true')
    parser.add_argument('--report', type=Path, default=ROOT / 'report_figures' / 'conversion_matrix_real_latest.json')
    args = parser.parse_args()
    artifacts_retained = bool(args.keep or args.workdir)
    context = None
    if args.workdir:
        workdir = args.workdir.expanduser().resolve(); workdir.mkdir(parents=True, exist_ok=True)
    elif args.keep:
        workdir = Path(tempfile.mkdtemp(prefix='streamdock-real-matrix-kept-'))
    else:
        context = tempfile.TemporaryDirectory(prefix='streamdock-real-matrix-'); workdir = Path(context.name)
    try:
        fixtures = build_all_fixtures(workdir)
        routes = []
        seen = set()
        for capability in list_capabilities():
            key = (capability.source, capability.target)
            if key in seen or capability.level == ConversionLevel.VENDOR or 'pdf' in key or key == ('pptx', 'png') or capability.source == 'folder':
                continue
            seen.add(key); routes.append(capability)
        expected_routes = set(json.loads(RELEASE_CONTRACT.read_text(encoding='utf-8'))['matrixRoutes'])
        actual_routes = {f'{item.source}->{item.target}' for item in routes}
        coverage_errors = {
            'missing': sorted(expected_routes - actual_routes),
            'unexpected': sorted(actual_routes - expected_routes),
        }
        results = []
        outputs = workdir / 'matrix-outputs'; outputs.mkdir()
        for index, capability in enumerate(routes):
            source, target = capability.source, capability.target
            started = time.perf_counter(); source_path = fixtures.get('json_object') if (source, target) == ('json', 'toml') else fixtures.get(source)
            row: dict[str, Any] = {'source': source, 'target': target, 'status': 'FAIL'}
            try:
                if source_path is None or not source_path.exists():
                    raise AssertionError(f'missing real fixture for {source}')
                route_output = outputs / f'{index:03d}-{source}-to-{target}'; route_output.mkdir()
                result = convert_file(source_path, source_path.name, source, target, route_output, image_quality=88)
                if not result.success or not result.output_path:
                    raise AssertionError(result.error or 'conversion produced no output')
                row.update(status='PASS', validation=validate_output(source, target, result.output_path), artifact=artifact_summary(result.output_path))
                if artifacts_retained:
                    row['outputPath'] = str(result.output_path)
            except Exception as exc:
                row['error'] = str(exc)
            row['elapsedSeconds'] = round(time.perf_counter() - started, 3); results.append(row)
        failed = [row for row in results if row['status'] != 'PASS']
        report = {'generatedAt': time.strftime('%Y-%m-%d %H:%M:%S'), 'workdir': str(workdir) if artifacts_retained else None, 'artifactsRetained': artifacts_retained, 'total': len(results), 'passed': len(results) - len(failed), 'failed': len(failed), 'coverage': {'expected': len(expected_routes), 'actual': len(actual_routes), **coverage_errors}, 'results': results}
        args.report.parent.mkdir(parents=True, exist_ok=True); args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print(f'REAL_CONVERSION_MATRIX={report["passed"]}/{report["total"]} passed')
        for row in failed:
            print(f'FAIL {row["source"]}->{row["target"]}: {row.get("error")}')
        if coverage_errors['missing'] or coverage_errors['unexpected']:
            print(f'FAIL route contract mismatch: {coverage_errors}')
        if args.keep or args.workdir:
            print(f'WORKDIR={workdir}')
        return 0 if not failed and not coverage_errors['missing'] and not coverage_errors['unexpected'] else 1
    finally:
        if context: context.cleanup()


if __name__ == '__main__':
    raise SystemExit(main())
