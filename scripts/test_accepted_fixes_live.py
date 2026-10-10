#!/usr/bin/env python3
"""Opt-in real-file HTTP acceptance. Isolated service/store; temp artifacts cleaned.
Requires FFmpeg/LibreOffice, openpyxl, python-docx, ebooklib and Playwright.
Never reads user samples, existing task state, cookies, or platform accounts.
"""
import base64
import csv
import io
import json
import os
import re
import socket
import subprocess
import sys
import tarfile
import tempfile
import time
import zipfile
from pathlib import Path
import requests
from PIL import Image, ImageDraw
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from docx import Document
from ebooklib import epub
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]

def main():
    with tempfile.TemporaryDirectory(prefix='streamdock-accepted-fixes-') as tmp:
        work = Path(tmp); outputs=work/'outputs'; outputs.mkdir()
        with socket.socket() as sock:
            sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
        base=f'http://127.0.0.1:{port}'
        env={**os.environ,'STREAMDOCK_MODE':'desktop','STREAMDOCK_TASK_STORAGE_PATH':str(work/'tasks.sqlite'),
             'STREAMDOCK_ARTIFACT_ROOT':str(work/'artifacts'),'TMPDIR':str(work),'PYTHONPATH':str(ROOT)}
        with (work/'service.log').open('w') as log:
            process=subprocess.Popen([sys.executable,'-m','uvicorn','app:app','--host','127.0.0.1','--port',str(port)],cwd=ROOT,env=env,stdout=log,stderr=log)
            count=0
            def convert(source,target):
                nonlocal count
                with source.open('rb') as f:
                    r=requests.post(base+'/api/convert/run',files={'file':(source.name,f)},data={'outputType':target,'outputPath':str(outputs)},timeout=180)
                r.raise_for_status();j=r.json();assert j.get('success'),j
                out=Path(j['outputPath']);assert out.is_file();count+=1
                print('REAL_FILE',source.suffix,'->',target,out.stat().st_size,flush=True)
                return out,j
            try:
                for _ in range(120):
                    try:
                        if requests.get(base+'/api/health/live',timeout=1).ok:break
                    except requests.RequestException:pass
                    if process.poll() is not None:raise RuntimeError((work/'service.log').read_text()[-3000:])
                    time.sleep(.2)
                # 120-row workbook: empty formatted first sheet, repeated/blank headers.
                w=Workbook();w.active['Z100'].number_format='0.00';sheet=w.create_sheet('中文数据')
                expected=[['姓名','','姓名']]+[[f'中文😀{i}',f'多行\n{i}',f'另一列{i}'] for i in range(120)]
                for row in expected:sheet.append(row)
                x=work/'report[Q1].xlsx';w.save(x)
                for target,delimiter in [('csv',','),('tsv','\t')]:
                    out,_=convert(x,target)
                    with out.open(encoding='utf-8',newline='') as f:assert list(csv.reader(f,delimiter=delimiter))==expected
                # Precise retry locates bracketed filenames; duplicate-key JSON rejected.
                with x.open('rb') as f:
                    failed=requests.post(base+'/api/convert/run',files={'file':(x.name,f)},data={'outputType':'json','outputPath':str(outputs)},timeout=120).json()
                assert not failed['success'] and '表头重复' in failed['error'],failed
                retry=requests.post(base+'/api/tasks/'+failed['task']['id']+'/retry',timeout=30)
                retried=retry.json()
                assert retry.status_code==200 and retried.get('mode')=='convert-retry',retry.text
                assert retried['task']['id']!=failed['task']['id'] and '表头重复' in retried.get('error',''),retry.text
                # Field order, literal formulas/errors and Unicode survive XLSX roundtrip.
                data=[{'z':'=1+1','a':'#N/A','中文':'😀'}];j=work/'literal.json';j.write_text(json.dumps(data,ensure_ascii=False))
                out,_=convert(j,'xlsx');wb=load_workbook(out);assert [c.value for c in wb.active[1]]==list(data[0]);assert all(c.data_type=='s' for c in wb.active[2]);wb.close()
                back,_=convert(out,'json');assert json.loads(back.read_text())==data
                # Arbitrary XML field names.
                fields={'业务 字段':'中文😀','1key':{'nested':'value'}};j=work/'keys.json';j.write_text(json.dumps(fields,ensure_ascii=False))
                out,_=convert(j,'xml');back,_=convert(out,'json');assert json.loads(back.read_text())['root']==fields
                # HTML boundaries and suppression of non-content elements.
                h=work/'content.html';h.write_text('<!doctype html><html><head><title>noise</title></head><body><p>第一段😀</p><table><tr><td>收入</td><td>120</td></tr></table><script>bad()</script></body></html>')
                out,_=convert(h,'txt');assert '收入\t120' in out.read_text() and 'noise' not in out.read_text()
                out,_=convert(h,'md');assert 'bad()' not in out.read_text() and '第一段😀' in out.read_text()
                # Source image, actual DOCX and RTF opened back through LibreOffice.
                image=work/'photo.png';im=Image.new('RGB',(160,100),'#1144aa');ImageDraw.Draw(im).rectangle((40,20,120,80),fill='#dd5511');im.save(image)
                doc=Document();doc.add_paragraph('文档中文😀');doc.add_picture(str(image));doc.add_table(rows=2,cols=2).cell(0,0).text='表格内容';source=work/'with-image.docx';doc.save(source)
                out,_=convert(source,'md');md=out.read_text();match=re.search(r'data:image/png;base64,([A-Za-z0-9+/=]+)',md);assert match and base64.b64decode(match[1])==image.read_bytes();assert '文档中文😀' in md
                out,_=convert(source,'rtf');assert out.read_bytes().startswith(b'{\\rtf') and b'\\pict' in out.read_bytes()
                from converters.adapters.document_basic import _libreoffice_convert
                readback=work/'rtf-readback';readback.mkdir();restored=_libreoffice_convert(out,readback,'docx')
                assert '文档中文😀' in '\n'.join(p.text for p in Document(restored).paragraphs)
                with zipfile.ZipFile(restored) as z:
                    pictures=[name for name in z.namelist() if name.startswith('word/media/')];assert pictures
                    with Image.open(io.BytesIO(z.read(pictures[0]))) as picture:picture.load();assert picture.size==(160,100)
                # EPUB chapter order and embedded image bytes, not only nonempty output.
                book=epub.EpubBook();book.set_identifier('fixture');book.set_title('中文书');book.set_language('zh')
                a=epub.EpubHtml(title='第一章',file_name='text/a.xhtml',lang='zh');a.content='<h1>第一章😀</h1><img src="../images/photo.png"/>'
                b=epub.EpubHtml(title='第二章',file_name='text/b.xhtml',lang='zh');b.content='<h1>第二章</h1><p>章节内容</p>'
                book.add_item(b);book.add_item(a);book.add_item(epub.EpubItem(uid='photo',file_name='images/photo.png',media_type='image/png',content=image.read_bytes()));book.spine=['a','b']
                book.add_item(epub.EpubNcx());book.add_item(epub.EpubNav());book.toc=[a,b]
                # EpubHtml identifiers default to chapter_<n>; use the actual IDs.
                book.spine=[a,b];e=work/'book.epub';epub.write_epub(str(e),book)
                for target in ('html','md','txt'):
                    out,_=convert(e,target);text=out.read_text();assert text.index('第一章')<text.index('第二章')
                    if target!='txt':
                        match=re.search(r'data:image/png;base64,([A-Za-z0-9+/=]+)',text);assert match and base64.b64decode(match[1])==image.read_bytes()
                # Archive roots, nested Unicode, empty directories, CRC and contents.
                z=work/'archive.zip';contents={'中文/文件😀.txt':'中文内容😀'.encode(),'report[Q1].txt':b'literal'}
                with zipfile.ZipFile(z,'w') as archive:
                    archive.writestr('empty/',b'')
                    for name,data in contents.items():archive.writestr(name,data)
                out,_=convert(z,'tar')
                with tarfile.open(out) as archive:
                    assert 'empty' in archive.getnames()
                    for name,data in contents.items():assert archive.extractfile(name).read()==data
                back,_=convert(out,'zip')
                with zipfile.ZipFile(back) as archive:
                    assert archive.testzip() is None
                    for name,data in contents.items():assert archive.read(name)==data
                # Folder TAR.GZ has the same root convention as ZIP.
                payload={'relativePaths':json.dumps(list(contents)), 'directoryPaths':json.dumps(['empty']), 'folderName':'root','outputType':'tar.gz','outputPath':str(outputs)}
                r=requests.post(base+'/api/convert/folder-run',files=[('files',(Path(n).name,io.BytesIO(d))) for n,d in contents.items()],data=payload,timeout=120);r.raise_for_status();j=r.json();assert j['success'],j
                with tarfile.open(j['outputPath']) as archive:
                    assert 'empty' in archive.getnames()
                    for name,data in contents.items():assert archive.extractfile(name).read()==data
                count+=1
                # SRT rolls over hours and is actually parsed.
                text=work/'long.txt';text.write_text('\n'.join(f'字幕中文😀{i}' for i in range(3661)))
                out,_=convert(text,'srt')
                from subtitles.service import parse_subtitles
                parsed=parse_subtitles(out.read_text(),filename=out.name);assert len(parsed.cues)==3661 and parsed.cues[-1].end==3661
                # Old workbook formats traverse LO -> XLSX -> common CSV logic.
                one=work/'single.xlsx';wb=Workbook();wb.active.append(['中文','字面']);wb.active.append(['😀','=1+1']);wb.active['B2'].data_type='s';wb.save(one)
                for old in ('xls','ods'):
                    old_dir=work/old;old_dir.mkdir();old_file=_libreoffice_convert(one,old_dir,old)
                    out,_=convert(old_file,'csv')
                    with out.open(encoding='utf-8',newline='') as f:assert list(csv.reader(f))==[['中文','字面'],['😀','=1+1']]
                # Real rendered responsive layout, no mocked fetch responses.
                with sync_playwright() as p:
                    browser=p.chromium.launch(headless=True)
                    for width in (390,1440):
                        page=browser.new_page(viewport={'width':width,'height':900});page.goto(base+'/use#parse')
                        assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                        box=page.locator('#outputType').bounding_box();assert box and box['width']>200
                        page.close()
                    browser.close()
                if '--bilibili-live' in sys.argv:
                    media=requests.post(base+'/api/fetch',json={
                        'link':'https://www.bilibili.com/video/BV1af4y1H7ga/',
                        'outputPath':str(outputs),'outputType':'mp4',
                        'saveAssets':True,'subtitleStrategy':'native'},timeout=180).json()
                    assert media.get('success') and '当前节点' in media.get('scopeWarning',''),media
                    video=media['outputPath']
                    probe=subprocess.run(['ffprobe','-v','error','-show_format','-show_streams','-of','json',video],capture_output=True,check=True,timeout=30)
                    info=json.loads(probe.stdout);assert 'mp4' in info['format']['format_name']
                    assert {'video','audio'} <= {s['codec_type'] for s in info['streams']}
                    subprocess.run(['ffmpeg','-v','error','-xerror','-i',video,'-f','null','-'],capture_output=True,check=True,timeout=60)
                    with Image.open(media['assets']['cover']) as cover:cover.load()
                    print('BILIBILI_REAL_HTTP=passed',media['scopeWarning'],info['format']['duration'],flush=True)
                print('ACCEPTED_FIXES_REAL_HTTP=passed',count,'successful conversions; bracket retry and protective rejection verified; mobile/desktop rendered',flush=True)
            finally:
                process.terminate()
                try:process.wait(timeout=15)
                except subprocess.TimeoutExpired:process.kill();process.wait()

if __name__=='__main__':main()
